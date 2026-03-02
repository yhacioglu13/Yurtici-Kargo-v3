"""Yurtiçi Kargo gönderilerinde maliyet kontrolü ve rapor üretimi."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.page import PageMargins

# BURASI ÖNEMLİ: config kökte, helpers app içinde
from config import MAIL_EXCEL_ROOT, TARIFF_EXCEL_PATH, RESULT_EXCEL_ROOT
from app.helpers import today_str_en, ensure_folder

def _ask_manual_excel_paths() -> list[Path]:
    """Kullanıcıdan 5. adım için bir veya birden fazla Excel dosyası seçmesini ister."""
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    selected_files = filedialog.askopenfilenames(
        title="5. Adım için maliyet kontrolünde kullanılacak Excel dosyalarını seçin",
        filetypes=[("Excel Dosyaları", "*.xlsx *.xls")],
    )
    root.destroy()

    paths = [Path(file) for file in selected_files]
    if not paths:
        print("⚠️ Dosya seçimi yapılmadı.")
        return []

    print(f"📂 Manuel seçim ile {len(paths)} dosya alındı.")
    for index, path in enumerate(paths, start=1):
        print(f"   {index}) {path}")
    return paths

def _load_and_merge_excels(paths: list[Path]) -> pd.DataFrame | None:
    """Verilen Excel dosyalarını tek DataFrame içinde birleştirir."""
    dataframes: list[pd.DataFrame] = []

    for path in paths:
        if not path.exists():
            print(f"⚠️ Dosya bulunamadı, atlanıyor: {path}")
            continue

        try:
            df = pd.read_excel(path)
            df["Kaynak Dosya"] = path.name
            dataframes.append(df)
            print(f"✅ Okundu: {path}")
        except Exception as exc:
            print(f"⚠️ Okunamadı, atlanıyor: {path} ({exc})")

    if not dataframes:
        print("❌ Geçerli bir Excel dosyası okunamadı.")
        return None

    return pd.concat(dataframes, ignore_index=True)
    
def load_base_data(manual_select: bool = False) -> pd.DataFrame | None:
    """Günlük birleştirilmiş mail excellerini okur."""
    if manual_select:
        selected_paths = _ask_manual_excel_paths()
        if not selected_paths:
            return None
        return _load_and_merge_excels(selected_paths)

    today = today_str_en()
    base_excel_path = Path(MAIL_EXCEL_ROOT) / today / f"{today} - Birleştirilmiş Mail Excelleri.xlsx"

    if not base_excel_path.exists():
        print(f"❌ Hata: Birleştirilmiş Excel bulunamadı! ({base_excel_path})")
        return None

    try:
        df = pd.read_excel(base_excel_path)
        print(f"📄 Birleştirilmiş Excel yüklendi: {base_excel_path}")
        return df
    except Exception as exc:
        print(f"❌ Hata: Birleştirilmiş Excel okunamadı! ({exc})")
        return None


def load_tariff() -> pd.DataFrame | None:
    """Tarife (fiyat) tablosunu okur."""
    tariff_path = Path(TARIFF_EXCEL_PATH)

    if not tariff_path.exists():
        print(f"❌ Hata: Tarife Excel dosyası bulunamadı! ({tariff_path})")
        return None

    try:
        df_tarife = pd.read_excel(tariff_path, sheet_name="Tarife")
        print(f"📄 Tarife Excel yüklendi: {tariff_path}")
        return df_tarife
    except Exception as exc:
        print(f"❌ Hata: Tarife Excel okunamadı! ({exc})")
        return None


def hesapla_maliyet_factory(df_tarife: pd.DataFrame):
    """Desi'ye göre maliyet hesaplayan fonksiyonu, tarife tablosuna göre üretir."""

    def get_cost_for_step(step: int) -> float:
        return df_tarife.loc[df_tarife["YURTİÇİ KARGO"] == step, "maliyet"].values[0]

    def hesapla_maliyet(desi) -> float | None:
        try:
            desi = float(desi)
        except Exception:
            return None

        if desi < 0:
            return None
        elif desi < 1:
            return get_cost_for_step(1)
        elif desi < 4:
            return get_cost_for_step(2)
        elif desi < 6:
            return get_cost_for_step(3)
        elif desi < 11:
            return get_cost_for_step(4)
        elif desi < 16:
            return get_cost_for_step(5)
        elif desi < 21:
            return get_cost_for_step(6)
        elif desi < 26:
            return get_cost_for_step(7)
        elif desi < 31:
            return get_cost_for_step(8)
        elif desi < 36:
            return get_cost_for_step(9)
        elif desi < 41:
            return get_cost_for_step(10)
        elif desi < 46:
            return get_cost_for_step(11)
        elif desi < 51:
            return get_cost_for_step(12)
        else:
            ana = get_cost_for_step(12)
            ek = get_cost_for_step(13)
            return ana + (desi - 50) * ek

    return hesapla_maliyet


def tutar_kontrol_et(row) -> str:
    """Ürün bedeli ile hesaplanan maliyeti karşılaştırır."""
    try:
        urun_bedeli = float(row["Ürün Bedeli"])
        maliyet = float(row["maliyet"])
    except Exception:
        return "Hesaplanamadı"

    if urun_bedeli == 0:
        return "Hesaplanamadı"

    fark_orani = abs(urun_bedeli - maliyet) / urun_bedeli
    return "Tutar uygun" if fark_orani <= 0.05 else "Tutar hatalı!!"


def filter_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Rapor için gerekli sütunları seçer."""
    columns_needed = [
        "Fatura Numarası",
        "Gönderici Müşteri",
        "Gönderen Müşteri Adresi",
        "Alıcı Müşteri",
        "Çıkış Birimi",
        "Varış Birimi",
        "Varış İli",
        "Kargo Tipi",
        "Toplam Kargo Adedi",
        "Desi / Kg",
        "Ürün Bedeli",
        "İrsaliye Matrahı",
        "Kdv",
        "İrsaliye Matrahı+KDV",
        "Mesafe (Km)",
        "Mesafe Açıklaması",
        "Teslim Alan",
        "maliyet",
        "Sonuç1",
    ]

    available_cols = [col for col in columns_needed if col in df.columns]
    missing = set(columns_needed) - set(available_cols)
    if missing:
        print(f"⚠️ Rapor için eksik sütunlar: {missing}")

    return df[available_cols].copy()


def apply_print_settings(sheet) -> None:
    """Sayfa yazdırma ayarlarını (print setup) düzenler.

    - A4, yatay (landscape)
    - Genişliği tek sayfaya sığdır
    - İlk satırı her sayfada tekrar et
    - Kenar boşlukları makul
    - Yatayda ortalanmış
    """
    # A4 + yatay
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4

    # Genişliği 1 sayfaya sığdır, yükseklik serbest
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    # Kenar boşlukları (inch cinsinden)
    sheet.page_margins = PageMargins(
        left=0.5,
        right=0.5,
        top=0.75,
        bottom=0.75,
        header=0.3,
        footer=0.3,
    )

    # Başlık satırını her sayfada tekrar et (1. satır)
    sheet.print_title_rows = "1:1"

    # Yatayda sayfayı ortala
    sheet.print_options.horizontalCentered = True


def save_per_invoice(df_filtered: pd.DataFrame) -> None:
    """Her fatura için ayrı Excel dosyaları oluşturur, biçimlendirir ve yazdırma ayarlarını yapar."""
    today = today_str_en()
    save_folder = ensure_folder(Path(RESULT_EXCEL_ROOT) / today)

    fatura_kolon_adi = "Fatura Numarası"
    if fatura_kolon_adi not in df_filtered.columns:
        print(f"❌ Hata: '{fatura_kolon_adi}' sütunu bulunamadı, fatura bazlı kayıt yapılamıyor.")
        return

    for fatura_no, grup in df_filtered.groupby(fatura_kolon_adi):
        dosya_adi = f"{fatura_no}.xlsx"
        dosya_yolu = save_folder / dosya_adi
        grup.to_excel(dosya_yolu, index=False)

        workbook = load_workbook(dosya_yolu)
        sheet = workbook.active

        # Hücrelerde metin sarma + hizalama
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    wrap_text=True,
                    horizontal="left",
                    vertical="top",
                )

        # İlk satır yüksekliği
        sheet.row_dimensions[1].height = 43.2

        # Sütun genişliklerini ayarla
        genislikler = {
            "A": 17,
            "B": 24,
            "C": 30,
            "D": 15,
            "E": 15,
            "F": 15,
            "G": 15,
            "H": 15,
            "I": 10,
            "J": 14,
            "K": 10,
            "L": 18,
            "M": 10,
            "N": 30,
            "O": 20,
            "P": 14,
            "Q": 15,
        }
        for sutun, genislik in genislikler.items():
            if sutun in sheet.column_dimensions:
                sheet.column_dimensions[sutun].width = genislik

        # 📄 Yazdırma ayarlarını uygula (A4, landscape, fit to width, header repeat)
        apply_print_settings(sheet)

        workbook.save(dosya_yolu)

    print(f"✅ Tüm fatura bazlı dosyalar başarıyla kaydedildi: {save_folder}")


def run_maliyet_kontrol(manual_select: bool = False) -> None:
    """Maliyet kontrolü yapar, özet tabloyu üretir ve fatura bazlı excelleri kaydeder."""
    df = load_base_data(manual_select=manual_select)
    if df is None:
        return

    df_tarife = load_tariff()
    if df_tarife is None:
        return

    # Maliyet hesapla
    hesapla_maliyet = hesapla_maliyet_factory(df_tarife)
    df["maliyet"] = df["Desi / Kg"].apply(hesapla_maliyet)

    # Tutar kontrolü
    df["Sonuç1"] = df.apply(tutar_kontrol_et, axis=1)

    # Filtrelenmiş tablo
    df_filtered = filter_columns(df)

    # Konsolda kısa bir ön izleme (var olan sütunlara göre)
    preview_cols = [
        c
        for c in ["Desi / Kg", "Ürün Bedeli", "maliyet", "Sonuç1"]
        if c in df_filtered.columns
    ]
    if preview_cols:
        print(df_filtered[preview_cols].head())
    else:
        print("⚠️ Önizleme için uygun sütun bulunamadı.")

    # Fatura bazlı dosyalar
    save_per_invoice(df_filtered)

